import io
import re
from datetime import date

import pandas as pd
import streamlit as st
import matplotlib.pyplot as plt
from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

st.set_page_config(page_title="Resume Intelligence System", page_icon="📄", layout="wide")

st.title("📄 Resume Intelligence System")
st.caption("Resume skill analysis • Job-role matching • Skill-gap insights • Learning roadmap")

JOB_ROLES = {
    "Data Analyst": [
        "python", "sql", "excel", "pandas", "numpy",
        "power bi", "statistics", "data visualization"
    ],
    "Machine Learning Engineer": [
        "python", "machine learning", "scikit-learn", "numpy",
        "pandas", "model evaluation", "fastapi", "docker"
    ],
    "AI Engineer": [
        "python", "deep learning", "llm", "apis",
        "machine learning", "pytorch", "tensorflow", "rag"
    ],
    "NLP Engineer": [
        "python", "nlp", "natural language processing",
        "transformers", "hugging face", "spacy", "pytorch"
    ],
    "Computer Vision Engineer": [
        "python", "opencv", "computer vision", "cnn",
        "deep learning", "pytorch", "tensorflow", "yolo"
    ],
    "Python Developer": [
        "python", "git", "apis", "fastapi",
        "flask", "django", "sql", "testing"
    ],
}

SKILL_CATEGORIES = {
    "Programming": ["python", "java", "c++", "c#", "javascript", "typescript", "r"],
    "Databases": ["sql", "mysql", "postgresql", "sqlite", "mongodb"],
    "Data Analysis": ["pandas", "numpy", "excel", "power bi", "tableau", "statistics", "data visualization"],
    "Machine Learning and AI": ["machine learning", "deep learning", "scikit-learn", "tensorflow", "pytorch", "llm", "rag", "model evaluation"],
    "NLP and Computer Vision": ["nlp", "natural language processing", "transformers", "hugging face", "spacy", "opencv", "computer vision", "cnn", "yolo"],
    "Web and Deployment": ["html", "css", "react", "fastapi", "flask", "django", "apis", "docker", "git", "testing", "aws", ".net"],
}

ALL_SKILLS = sorted(
    {skill for skills in SKILL_CATEGORIES.values() for skill in skills},
    key=len,
    reverse=True,
)

ROADMAP = {
    "python": "Practise Python fundamentals, functions, data structures, and small projects.",
    "sql": "Practise SELECT, JOIN, GROUP BY, subqueries, and query challenges.",
    "pandas": "Work through data cleaning, grouping, merging, and exploratory analysis.",
    "numpy": "Practise arrays, vectorized operations, indexing, and numerical methods.",
    "statistics": "Review distributions, descriptive statistics, probability, and hypothesis testing.",
    "power bi": "Build a dashboard with data modelling, measures, and visual storytelling.",
    "data visualization": "Create clear charts and explain the insight each chart communicates.",
    "machine learning": "Study supervised learning, preprocessing, model selection, and evaluation.",
    "scikit-learn": "Build end-to-end pipelines with train/test splits and cross-validation.",
    "model evaluation": "Practise suitable metrics, validation, error analysis, and baseline comparison.",
    "fastapi": "Create a small REST API and expose a model or function through endpoints.",
    "docker": "Containerize a small Python application and learn images, containers, and ports.",
    "deep learning": "Learn neural-network fundamentals and train a small model.",
    "pytorch": "Complete a tensor, training-loop, and model-evaluation tutorial.",
    "tensorflow": "Build and evaluate a small neural network.",
    "llm": "Learn LLM basics, prompting, limitations, and evaluation.",
    "rag": "Build a small retrieval-augmented generation prototype and evaluate retrieval quality.",
    "apis": "Practise HTTP methods, JSON, request/response handling, and API errors.",
    "nlp": "Study tokenization, text normalization, vectorization, and text classification.",
    "natural language processing": "Implement text preprocessing and a small NLP task.",
    "transformers": "Learn attention and transformer architecture; fine-tune or use a small model.",
    "hugging face": "Use the Transformers pipeline and model/tokenizer APIs in a small demo.",
    "spacy": "Practise spaCy tokenization, entities, and phrase matching.",
    "opencv": "Build an image-processing demo using OpenCV.",
    "computer vision": "Study image preprocessing, classification, and evaluation.",
    "cnn": "Train or inspect a small convolutional neural network.",
    "yolo": "Run a supported object-detection demo and review precision/recall.",
    "git": "Practise commits, branches, merges, and a clean GitHub workflow.",
    "testing": "Write unit tests for valid inputs, invalid files, and edge cases.",
    "flask": "Build a small Flask app with routes and input validation.",
    "django": "Create a small Django project with a model, view, and template.",
    "aws": "Learn core cloud concepts and deploy a small non-sensitive demo.",
    "java": "Practise OOP, collections, exception handling, and a small project.",
    "javascript": "Build interactive web components and practise asynchronous requests.",
    "typescript": "Practise types, interfaces, and typed functions in a small application.",
    "mysql": "Practise relational schema design, joins, indexes, and constraints.",
    "postgresql": "Practise SQL queries, schema design, and database constraints.",
    "sqlite": "Build a small Python app that reads and writes a SQLite database.",
    "mongodb": "Practise document modelling and CRUD operations.",
    "excel": "Practise formulas, pivot tables, cleaning, and charts.",
    "tableau": "Create a dashboard with filters and explain the main findings.",
    "html": "Build a semantic multi-section webpage.",
    "css": "Practise responsive layouts, flexbox/grid, and accessible styling.",
    "react": "Build a small component-based UI with state and props.",
    "c++": "Practise core syntax, OOP, STL, and problem-solving exercises.",
    "c#": "Practise C# syntax, OOP, and a small application.",
    "r": "Practise data frames, visualization, and basic statistical analysis.",
    ".net": "Build a small .NET application and understand its project structure.",
}

def extract_resume_text(uploaded_file):
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        table_text = [
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        ]
        text = "\n".join(paragraphs + table_text)
    else:
        raise ValueError("Please upload a PDF or DOCX resume.")
    if not text.strip():
        raise ValueError("No text could be extracted. If this is a scanned PDF, use a text-based PDF or add OCR.")
    return text

def clean_text(text):
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[^a-z0-9+#.\s/-]", " ", text)
    return re.sub(r"\s+", " ", text).strip()

def skill_pattern(skill):
    # Special cases for symbols/punctuation used in technical names.
    if skill == "c++":
        return r"(?<![a-z0-9])c\+\+(?![a-z0-9])"
    if skill == "c#":
        return r"(?<![a-z0-9])c#(?![a-z0-9])"
    if skill == ".net":
        return r"(?<![a-z0-9])\.net(?![a-z0-9])"
    return r"(?<![a-z0-9])" + re.escape(skill) + r"(?![a-z0-9])"

def extract_skills(text):
    normalized = clean_text(text)
    found = []
    for skill in ALL_SKILLS:
        if re.search(skill_pattern(skill), normalized):
            found.append(skill)
    return set(found)

def score_roles(resume_text):
    role_names = list(JOB_ROLES)
    role_docs = [" ".join(JOB_ROLES[role]) for role in role_names]
    documents = [clean_text(resume_text)] + [clean_text(doc) for doc in role_docs]
    vectorizer = TfidfVectorizer(ngram_range=(1, 2))
    matrix = vectorizer.fit_transform(documents)
    scores = cosine_similarity(matrix[0:1], matrix[1:]).flatten()
    frame = pd.DataFrame({"Job Role": role_names, "Match Score (%)": (scores * 100).round(1)})
    return frame.sort_values("Match Score (%)", ascending=False).reset_index(drop=True)

def build_roadmap(missing_skills):
    if not missing_skills:
        return ["No gaps were identified by the current skill dictionary. Review the role requirements manually."]
    return [f"{i}. {skill.title()}: {ROADMAP.get(skill, 'Study the fundamentals, complete a small practical exercise, and add the result to a portfolio project.')}"
            for i, skill in enumerate(sorted(missing_skills), start=1)]

with st.sidebar:
    st.header("How it works")
    st.write("1. Upload a resume")
    st.write("2. Extract text and detect listed skills")
    st.write("3. Compare text with six example roles")
    st.write("4. Review skill gaps and learning suggestions")
    st.divider()
    st.caption("Privacy: this app processes the uploaded file during the session. Do not upload a resume unless you are comfortable processing it in this environment.")
    st.caption("Scores are text-similarity estimates, not hiring probabilities.")

uploaded_file = st.file_uploader("Upload your resume (PDF or DOCX)", type=["pdf", "docx"])

if uploaded_file:
    try:
        resume_text = extract_resume_text(uploaded_file)
        skills_found = extract_skills(resume_text)
        results = score_roles(resume_text)

        st.success(f"Resume parsed: {uploaded_file.name} · {len(resume_text):,} characters extracted")
        with st.expander("Preview extracted text"):
            st.text(resume_text[:4000])

        st.subheader("Detected skills")
        if skills_found:
            categorized = {}
            for category, skills in SKILL_CATEGORIES.items():
                matches = sorted(set(skills) & skills_found)
                if matches:
                    categorized[category] = matches
            cols = st.columns(2)
            for idx, (category, skills) in enumerate(categorized.items()):
                with cols[idx % 2]:
                    st.markdown(f"**{category}**")
                    st.write(", ".join(s.title() for s in skills))
        else:
            st.warning("No listed skills were detected. Check the extracted text and expand the skill dictionary if needed.")

        st.subheader("Recommended job roles")
        st.dataframe(results, use_container_width=True, hide_index=True)
        chart = results.sort_values("Match Score (%)", ascending=True)
        fig, ax = plt.subplots(figsize=(9, 4))
        ax.barh(chart["Job Role"], chart["Match Score (%)"])
        ax.set_xlabel("TF-IDF cosine similarity (%)")
        ax.set_xlim(0, max(10, float(chart["Match Score (%)"].max()) * 1.15))
        ax.set_title("Resume-to-role text similarity")
        fig.tight_layout()
        st.pyplot(fig)
        plt.close(fig)

        st.markdown("#### Top three roles")
        for i, row in results.head(3).iterrows():
            st.write(f"{i+1}. **{row['Job Role']}** — {row['Match Score (%)']}% text similarity")

        st.subheader("Skill-gap analysis")
        selected_role = st.selectbox("Choose a target role", list(JOB_ROLES.keys()), index=list(JOB_ROLES.keys()).index(results.iloc[0]["Job Role"]))
        required = set(JOB_ROLES[selected_role])
        matched = sorted(required & skills_found)
        missing = sorted(required - skills_found)

        c1, c2, c3 = st.columns(3)
        c1.metric("Required skills", len(required))
        c2.metric("Skills found", len(matched))
        c3.metric("Skills missing", len(missing))

        st.markdown("**Skills found for this role**")
        st.write(", ".join(s.title() for s in matched) if matched else "No required skills from this role were detected.")
        st.markdown("**Skills to develop**")
        st.write(", ".join(s.title() for s in missing) if missing else "No missing skills detected by this dictionary.")

        st.subheader("Suggested learning roadmap")
        for item in build_roadmap(missing):
            st.write(item)

        report_rows = []
        for _, row in results.iterrows():
            role = row["Job Role"]
            req = set(JOB_ROLES[role])
            report_rows.append({
                "Job Role": role,
                "Match Score (%)": row["Match Score (%)"],
                "Skills Found": ", ".join(sorted(req & skills_found)),
                "Missing Skills": ", ".join(sorted(req - skills_found)),
            })
        report_df = pd.DataFrame(report_rows)
        csv_bytes = report_df.to_csv(index=False).encode("utf-8")
        st.download_button(
            "⬇️ Download analysis report (CSV)",
            data=csv_bytes,
            file_name="resume_intelligence_report.csv",
            mime="text/csv",
        )

        text_report = [
            "RESUME INTELLIGENCE SYSTEM — ANALYSIS REPORT",
            f"Generated: {date.today().isoformat()}",
            f"Resume file: {uploaded_file.name}",
            f"Extracted characters: {len(resume_text)}",
            f"Detected skills: {', '.join(sorted(skills_found)) or 'None'}",
            "",
            "ROLE RANKING:",
        ]
        text_report += [f"{i+1}. {r['Job Role']} — {r['Match Score (%)']}% text similarity" for i, r in results.iterrows()]
        text_report += ["", f"TARGET ROLE: {selected_role}", f"Skills found: {', '.join(matched) or 'None'}",
                        f"Missing skills: {', '.join(missing) or 'None'}", "", "LEARNING ROADMAP:"]
        text_report += build_roadmap(missing)
        st.download_button(
            "⬇️ Download readable report (TXT)",
            data="\n".join(text_report).encode("utf-8"),
            file_name="resume_intelligence_report.txt",
            mime="text/plain",
        )

        st.info("This prototype uses a manually curated skill list and six example job profiles. Results depend on the wording of the resume and role descriptions; they should not be treated as an employment decision.")
    except Exception as exc:
        st.error(f"Could not analyze this resume: {exc}")
else:
    st.info("Upload a PDF or DOCX resume to begin.")
    st.markdown("### Project features")
    st.markdown("- PDF/DOCX text extraction\n- Categorized keyword-based skill detection\n- TF-IDF and cosine-similarity role ranking\n- Top-three recommendations\n- Target-role skill gaps\n- Rule-based learning roadmap\n- Downloadable CSV and TXT reports")
